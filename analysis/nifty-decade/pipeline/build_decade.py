"""Decade workbook: Nifty 50 level, NSE trailing P/E and trailing earnings, Sep 2016 to Sep 2026.

Inputs (blue): daily Nifty closes (Yahoo ^NSEI, double-transcribed), NSE month-end P/E (Downstox republication of
NSE Indices data, cross-checked), the standalone-to-consolidated switch-day P/Es. Everything else is a formula.
"""
import sys, glob, json
import datetime as dt
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import LineChart, ScatterChart, Reference, Series
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.label import DataLabelList, DataLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont

SCR = '/tmp/claude-0/-home-user-Blockdeals/79809461-fe1d-550f-bdd0-e82c580f77b9/scratchpad'
OUT = f'{SCR}/decade/nifty_decade_price_pe_earnings.xlsx'
LOGO = ('/root/.claude/skills/synced/5df40f86-7c7f-40a9-b3c9-2f700a72c2cf_4c4812c9-5adb-4ea7-8980-b5669d92109d/'
        '1finance-brand-guidelines/assets/1finance-logo-vertical-black-200.png')
CFG = json.load(open(f'{SCR}/decade/inputs.json'))   # switch-day values, latest P/E, sources

# brand tokens
C_PRICE = '6D57AA'                      # purple 700
C_EPS, C_EPS_SA = 'CF6827', 'E49B53'    # orange 600 / orange 400
C_PE, C_PE_SA = '356494', '6A9AC6'      # blue 600 / blue 400
C_REF, C_GRID, C_AXIS = '8E8C81', 'E4E4E0', 'BCBCB5'
INK, INK2, INK3 = '272623', '6C6A62', '8E8C81'
BRAND, BODY, BLUE_INPUT = 'Fira Sans', 'Arial', '0000FF'
thin = Side(style='thin', color='D4D4CF')
hdr_fill = PatternFill('solid', fgColor='F6F6FC')

# ---------------------------------------------------------------- inputs
px = pd.concat([pd.read_csv(p) for p in sorted(glob.glob(f'{SCR}/data/final/NSEI_D*.csv'))])
px['date'] = pd.to_datetime(px['date'])
px = px.drop_duplicates('date').sort_values('date')
px = px[(px['date'] >= '2016-09-01') & (px['date'] <= '2026-09-28')].reset_index(drop=True)
flat = px['close'].eq(px['close'].shift()) & px['volume'].isna()
px = px[~flat].reset_index(drop=True)

pe = pd.read_csv(f'{SCR}/decade/nse_pe_monthly_downstox.csv')
pe = pe[(pe.y * 100 + pe.m >= 201609) & (pe.y * 100 + pe.m <= 202609)].reset_index(drop=True)
me = px.groupby([px.date.dt.year, px.date.dt.month])['date'].max().reset_index(drop=True)
me = pd.DataFrame({'date': me, 'y': me.dt.year, 'm': me.dt.month})
mon = pe.merge(me, on=['y', 'm'], how='left')
assert mon['date'].notna().all(), mon[mon['date'].isna()]
# September 2026: use the latest published reading (see inputs.json) instead of the mid-month value
mon.loc[(mon.y == 2026) & (mon.m == 9), 'pe'] = CFG['latest_pe']
mon['pe_note'] = ''
mon.loc[(mon.y == 2026) & (mon.m == 9), 'pe_note'] = CFG['latest_pe_note']

wb = Workbook()
ws_c1 = wb.active; ws_c1.title = 'Chart'
ws_m = wb.create_sheet('Monthly')
ws_d = wb.create_sheet('Daily')
ws_s = wb.create_sheet('Summary')
ws_v = wb.create_sheet('Validation')
ws_x = wb.create_sheet('Method & Sources')


def hdr(ws, row, labels, widths=None, height=None):
    for j, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = Border(bottom=thin)
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    if height:
        ws.row_dimensions[row].height = height


def f(ws, cell, value, fmt=None, color='000000', bold=False, size=10, italic=False):
    c = ws[cell] if isinstance(cell, str) else ws.cell(row=cell[0], column=cell[1])
    c.value = value
    if fmt: c.number_format = fmt
    c.font = Font(name=BODY, size=size, color=color, bold=bold, italic=italic)
    c.alignment = Alignment(wrap_text=False, vertical='center')
    return c


# ================================================================= Daily
hdr(ws_d, 1, ['Date', 'Nifty 50 close', 'Note'], [12, 14, 60], 30)
ws_d.freeze_panes = 'A2'
D0 = 2
for i, r in enumerate(px.itertuples(index=False)):
    f(ws_d, (D0 + i, 1), r.date.date(), 'dd-mmm-yyyy')
    f(ws_d, (D0 + i, 2), float(r.close), '#,##0.00', BLUE_INPUT)
    if r.date.strftime('%Y-%m-%d') == '2019-03-29':
        f(ws_d, (D0 + i, 3), "Missing from Yahoo's daily table; value is Yahoo's March 2019 monthly close.", None, INK2, size=9)
D1 = D0 + len(px) - 1

# ================================================================= Summary: inputs referenced by Monthly
ws = ws_s
ws.column_dimensions['A'].width = 52
for col, w in zip('BCDEFGHIJK', [13, 12, 12, 13, 13, 13, 13, 13, 13, 13]):
    ws.column_dimensions[col].width = w
f(ws, 'A1', 'Nifty 50 over ten years: price, NSE P/E and earnings', bold=True, size=13, color=INK)
f(ws, 'A2', 'Blue = input. Everything else is a formula on the Monthly and Daily sheets. Price = earnings x P/E, so each row reconciles.',
  italic=True, size=9, color=INK2)
f(ws, 'A4', "NSE's switch from standalone to consolidated earnings", bold=True, color=INK)
f(ws, 'A5', 'Nifty close, ' + CFG['switch_prev_date'] + ' (last standalone day)'); f(ws, 'B5', CFG['switch_prev_close'], '#,##0.00', BLUE_INPUT)
f(ws, 'A6', 'NSE P/E that day, standalone earnings'); f(ws, 'B6', CFG['switch_prev_pe'], '0.00', BLUE_INPUT)
f(ws, 'A7', 'Nifty close, ' + CFG['switch_date'] + ' (first consolidated day)'); f(ws, 'B7', CFG['switch_close'], '#,##0.00', BLUE_INPUT)
f(ws, 'A8', 'NSE P/E that day, consolidated earnings'); f(ws, 'B8', CFG['switch_pe'], '0.00', BLUE_INPUT)
f(ws, 'A9', 'Consolidated / standalone earnings at the switch', bold=True); f(ws, 'B9', '=(B7/B8)/(B5/B6)', '0.000', bold=True)
f(ws, 'A10', 'Consolidated / standalone, Bloomberg 10-year average (input)'); f(ws, 'B10', CFG['bbg_avg_ratio'], '0.000', BLUE_INPUT)
f(ws, 'C10', 'The Hindu BusinessLine, Mar 2021: "on the average ... 14 per cent higher"', None, INK2, size=9)
f(ws, 'A11', 'Ratio used for the like-for-like columns (input)', bold=True); f(ws, 'B11', '=B9', '0.000', BLUE_INPUT, bold=True)
f(ws, 'C11', 'Default: the switch-day ratio. Type 1.14 (or anything else) to see how the like-for-like columns move.', None, INK2, size=9)
f(ws, 'A12', 'Low estimate, PrimeInvestor Dec 2019 (input)'); f(ws, 'B12', CFG['low_ratio'], '0.000', BLUE_INPUT)
f(ws, 'C12', 'PrimeInvestor: consolidated earnings "about 5 per cent higher than the standalone earnings"', None, INK2, size=9)

# ================================================================= Monthly
ws = ws_m
cols = ['Month-end (last trading day)', 'Nifty 50 close', 'NSE P/E as published', 'Earnings basis', 'Trailing EPS as published (close / P/E), ₹',
        'Trailing EPS, like-for-like (pre-2021 x ratio), ₹', 'P/E, like-for-like (pre-2021 / ratio)',
        'EPS growth over 12 months, like-for-like', 'Note']
hdr(ws, 1, cols, [12, 12, 11, 13, 14, 14, 13, 12, 70], 58)
ws.freeze_panes = 'B2'
M0 = 2
M1 = M0 + len(mon) - 1
for i, r in enumerate(mon.itertuples(index=False)):
    row = M0 + i
    f(ws, (row, 1), r.date.date(), 'mmm-yyyy', BLUE_INPUT)
    f(ws, (row, 2), f'=INDEX(Daily!$B:$B,MATCH(A{row},Daily!$A:$A,0))', '#,##0')
    f(ws, (row, 3), float(r.pe), '0.0"x"', BLUE_INPUT)
    f(ws, (row, 4), r.basis, None, BLUE_INPUT)
    f(ws, (row, 5), f'=B{row}/C{row}', '"₹"#,##0')
    f(ws, (row, 6), f'=IF(D{row}="standalone",E{row}*Summary!$B$11,E{row})', '#,##0.0')
    f(ws, (row, 7), f'=B{row}/F{row}', '0.0')
    if i >= 12:
        f(ws, (row, 8), f'=F{row}/F{row-12}-1', '0%')
    if r.pe_note:
        f(ws, (row, 9), r.pe_note, None, INK2, size=9)
sw_row = M0 + int(((mon.y * 100 + mon.m) < 202103).sum())
ws.cell(row=sw_row, column=9, value=('First consolidated reading (31 Mar 2021). NSE P/E was 40.43 the day before on standalone earnings; '
                                      'the index barely moved.')).font = Font(name=BODY, size=9, color=INK2)

# ================================================================= Summary tables
ws = ws_s
hdr_row = 14
labels = ['Point in time', 'Month-end', 'Nifty 50', 'NSE P/E as published', 'Basis', 'EPS as published (₹)',
          'EPS like-for-like (₹)', 'P/E like-for-like', 'Price multiple to latest', 'EPS multiple to latest (like-for-like)',
          'P/E multiple to latest (like-for-like)']
for j, lab in enumerate(labels, start=1):
    c = ws.cell(row=hdr_row, column=j, value=lab); c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[hdr_row].height = 44
points = [
    ('Ten years ago', f'=INDEX(Monthly!$A:$A,{M0})'),
    ('End of 2019 (before COVID)', '=INDEX(Monthly!$A:$A,MATCH(DATE(2019,12,31),Monthly!$A:$A,1))'),
    ('March 2020 (COVID crash)', '=INDEX(Monthly!$A:$A,MATCH(DATE(2020,3,31),Monthly!$A:$A,1))'),
    ('Feb 2021 (last standalone month-end)', '=INDEX(Monthly!$A:$A,MATCH(DATE(2021,2,28),Monthly!$A:$A,1))'),
    ('Mar 2021 (first consolidated reading)', '=INDEX(Monthly!$A:$A,MATCH(DATE(2021,3,31),Monthly!$A:$A,1))'),
    ('Sep 2024 (record Nifty high month)', '=INDEX(Monthly!$A:$A,MATCH(DATE(2024,9,30),Monthly!$A:$A,1))'),
    ('End of 2025', '=INDEX(Monthly!$A:$A,MATCH(DATE(2025,12,31),Monthly!$A:$A,1))'),
    ('Latest (28 Sep 2026)', f'=INDEX(Monthly!$A:$A,{M1})'),
]
P0 = hdr_row + 1
LATEST = P0 + len(points) - 1
for k, (lab, date_f) in enumerate(points):
    r = P0 + k
    f(ws, (r, 1), lab, None, INK, bold=(k == len(points) - 1))
    f(ws, (r, 2), date_f, 'mmm-yyyy')
    m = f'MATCH($B{r},Monthly!$A:$A,0)'
    f(ws, (r, 3), f'=INDEX(Monthly!$B:$B,{m})', '#,##0')
    f(ws, (r, 4), f'=INDEX(Monthly!$C:$C,{m})', '0.0"x"')
    f(ws, (r, 5), f'=INDEX(Monthly!$D:$D,{m})')
    f(ws, (r, 6), f'=INDEX(Monthly!$E:$E,{m})', '"₹"#,##0')
    f(ws, (r, 7), f'=INDEX(Monthly!$F:$F,{m})', '"₹"#,##0')
    f(ws, (r, 8), f'=INDEX(Monthly!$G:$G,{m})', '0.0"x"')
    f(ws, (r, 9), f'=$C${LATEST}/C{r}', '0.00"x"')
    f(ws, (r, 10), f'=$G${LATEST}/G{r}', '0.00"x"')
    f(ws, (r, 11), f'=$H${LATEST}/H{r}', '0.00"x"')
ws.cell(row=LATEST, column=1).fill = hdr_fill

r = LATEST + 2
f(ws, (r, 1), 'Two halves of the decade (as published; the second half crosses the 2021 switch, see the like-for-like range below)', bold=True, color=INK)
hdr2 = ['Period', 'From', 'To', 'Price change', 'EPS change as published', 'P/E change as published', 'Price CAGR']
for j, lab in enumerate(hdr2, start=1):
    c = ws.cell(row=r + 1, column=j, value=lab); c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True)
ws.row_dimensions[r + 1].height = 30
halves = [('Sep 2016 to Dec 2019 (both standalone: like-for-like)', P0, P0 + 1),
          ('Dec 2019 to latest', P0 + 1, LATEST), ('Whole decade', P0, LATEST)]
H0 = r + 2
for k, (lab, a, b) in enumerate(halves):
    rr = H0 + k
    f(ws, (rr, 1), lab, None, INK)
    f(ws, (rr, 2), f'=B{a}', 'mmm-yyyy'); f(ws, (rr, 3), f'=B{b}', 'mmm-yyyy')
    f(ws, (rr, 4), f'=C{b}/C{a}-1', '+0%;-0%')
    f(ws, (rr, 5), f'=F{b}/F{a}-1', '+0%;-0%')
    f(ws, (rr, 6), f'=D{b}/D{a}-1', '+0%;-0%')
    f(ws, (rr, 7), f'=(C{b}/C{a})^(365.25/(B{b}-B{a}))-1', '0.0%')

S0 = H0 + 5
f(ws, (S0, 1), 'Like-for-like range: pre-2021 earnings restated to a consolidated basis', bold=True, color=INK)
for j, lab in enumerate(['Measure', 'Ratio 1.05 (PrimeInvestor, low)', 'Ratio 1.14 (Bloomberg 10-yr average)', 'Switch-day ratio'], start=1):
    c = ws.cell(row=S0 + 1, column=j, value=lab); c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True)
ws.row_dimensions[S0 + 1].height = 30
f(ws, (S0 + 2, 1), 'Ratio'); f(ws, (S0 + 2, 2), '=$B$12', '0.000'); f(ws, (S0 + 2, 3), '=$B$10', '0.000'); f(ws, (S0 + 2, 4), '=$B$9', '0.000')
sens = [
    ('Earnings multiple over the decade', f'=$F${LATEST}/($F${P0}*{{R}})', '0.00"x"'),
    ('P/E ten years ago, restated', f'=$D${P0}/{{R}}', '0.0"x"'),
    ('Change in P/E over the decade', f'=$D${LATEST}/($D${P0}/{{R}})-1', '+0%;-0%'),
    ('Earnings multiple, Dec 2019 to latest', f'=$F${LATEST}/($F${P0 + 1}*{{R}})', '0.00"x"'),
    ('P/E at end-2019, restated', f'=$D${P0 + 1}/{{R}}', '0.0"x"'),
    ('Change in P/E, Dec 2019 to latest', f'=$D${LATEST}/($D${P0 + 1}/{{R}})-1', '+0%;-0%'),
    ('P/E in March 2020, restated', f'=$D${P0 + 2}/{{R}}', '0.0"x"'),
]
for k, (lab, formula, fmt) in enumerate(sens):
    rr = S0 + 3 + k
    f(ws, (rr, 1), lab)
    f(ws, (rr, 2), formula.replace('{R}', f'$B${S0 + 2}'), fmt)
    f(ws, (rr, 3), formula.replace('{R}', f'$C${S0 + 2}'), fmt)
    f(ws, (rr, 4), formula.replace('{R}', f'$D${S0 + 2}'), fmt)
SENS = {'eps_mult': S0 + 3, 'pe_then': S0 + 4, 'pe_chg': S0 + 5, 'eps_mult_h2': S0 + 6, 'pe_2019': S0 + 7, 'pe_chg_h2': S0 + 8, 'pe_mar20': S0 + 9}

Q0 = S0 + 3 + len(sens) + 1
f(ws, (Q0, 1), "Where today's P/E sits", bold=True, color=INK)
f(ws, (Q0 + 1, 1), 'Lowest month-end NSE P/E since the switch (Mar 2021 onward)')
f(ws, (Q0 + 1, 2), f'=MIN(Monthly!$C${sw_row}:$C${M1})', '0.00"x"')
f(ws, (Q0 + 1, 3), f'=INDEX(Monthly!$A${sw_row}:$A${M1},MATCH(B{Q0 + 1},Monthly!$C${sw_row}:$C${M1},0))', 'mmm-yyyy')
f(ws, (Q0 + 2, 1), 'Median month-end NSE P/E since the switch')
f(ws, (Q0 + 2, 2), f'=MEDIAN(Monthly!$C${sw_row}:$C${M1})', '0.0"x"')
f(ws, (Q0 + 3, 1), 'Second-lowest month-end reading since the switch')
f(ws, (Q0 + 3, 2), f'=SMALL(Monthly!$C${sw_row}:$C${M1},2)', '0.00"x"')
f(ws, (Q0 + 3, 3), f'=INDEX(Monthly!$A${sw_row}:$A${M1},MATCH(B{Q0 + 3},Monthly!$C${sw_row}:$C${M1},0))', 'mmm-yyyy')
f(ws, (Q0 + 4, 1), 'Median month-end NSE P/E, Sep 2016 to Feb 2021 (standalone, as published)')
f(ws, (Q0 + 4, 2), f'=MEDIAN(Monthly!$C${M0}:$C${sw_row - 1})', '0.0"x"')
f(ws, (Q0 + 5, 1), '... restated at the ratio in B11')
f(ws, (Q0 + 5, 2), f'=B{Q0 + 4}/$B$11', '0.0"x"')
SUM_ROWS = dict(P0=P0, LATEST=LATEST, H0=H0, S0=S0, Q0=Q0)

# ================================================================= Validation
ws = ws_v
for col, w in zip('ABCDEFG', [14, 14, 13, 11, 52, 18, 80]):
    ws.column_dimensions[col].width = w
f(ws, 'A1', 'A. NSE P/E: Downstox values checked against NSE\'s own Index Dashboards and Nippon India notes', bold=True, size=12, color=INK)
for j, lab in enumerate(['Date', 'Downstox (used)', 'Second source', 'Difference', 'Second source', 'Type', 'URL'], start=1):
    c = ws.cell(row=2, column=j, value=lab); c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
for k, v in enumerate(CFG['crosscheck']):
    r = 3 + k
    f(ws, (r, 1), dt.date.fromisoformat(v['date']), 'dd-mmm-yyyy')
    f(ws, (r, 2), v['downstox'], '0.00', BLUE_INPUT); f(ws, (r, 3), v['other'], '0.00', BLUE_INPUT)
    f(ws, (r, 4), f'=ROUND(C{r}-B{r},2)', '+0.00;-0.00;0.00')
    f(ws, (r, 5), v['source']); f(ws, (r, 6), v.get('type', '')); f(ws, (r, 7), v.get('url', ''))
rA = 3 + len(CFG['crosscheck'])
f(ws, (rA, 1), 'Largest absolute difference'); f(ws, (rA, 4), f'=SUMPRODUCT(MAX(ABS(D3:D{rA - 1})))', '0.00')
f(ws, (rA + 1, 1), 'Downstox shows one decimal except at year-ends, so differences up to 0.05 are rounding.', italic=True, size=9, color=INK2)

rv = rA + 3
f(ws, (rv, 1), 'B. How much bigger were consolidated earnings than standalone? (the ratio used to restate pre-2021 numbers)', bold=True, size=12, color=INK)
for j, lab in enumerate(['When', 'Ratio', '', '', 'Source', '', 'URL'], start=1):
    c = ws.cell(row=rv + 1, column=j, value=lab); c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
for k, v in enumerate(CFG['ratio_evidence']):
    r = rv + 2 + k
    f(ws, (r, 1), v['when']); f(ws, (r, 2), v['ratio'], '0.000', BLUE_INPUT); f(ws, (r, 5), v['source']); f(ws, (r, 7), v.get('url', ''))
ws.cell(row=rv + 2, column=2).value = '=Summary!B9'
ws.cell(row=rv + 2, column=2).font = Font(name=BODY, size=10, color='000000')

rv2 = rv + 3 + len(CFG['ratio_evidence'])
f(ws, (rv2, 1), 'C. Actual Nifty EPS tallied by Motilal Oswal (fiscal years ending March) vs NSE-implied trailing EPS two months after year-end', bold=True, size=12, color=INK)
for j, lab in enumerate(['Fiscal year', 'MOFSL EPS (₹)', 'NSE-implied EPS end-May, as published (₹)', 'NSE / MOFSL', 'Note / source'], start=1):
    c = ws.cell(row=rv2 + 1, column=j, value=lab); c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True)
ws.row_dimensions[rv2 + 1].height = 30
for k, v in enumerate(CFG['mofsl_eps']):
    r = rv2 + 2 + k
    f(ws, (r, 1), v['fy']); f(ws, (r, 2), v['eps'], '#,##0', BLUE_INPUT)
    yr = int(v['fy'][2:]) + 2000
    if yr >= 2017:
        f(ws, (r, 3), f'=INDEX(Monthly!$E:$E,MATCH(DATE({yr},5,31),Monthly!$A:$A,1))', '#,##0')
        f(ws, (r, 4), f'=C{r}/B{r}', '0.00')
    f(ws, (r, 5), v.get('note', ''))
re_ = rv2 + 2 + len(CFG['mofsl_eps'])
f(ws, (re_, 1), 'MOFSL FY16 to FY26 EPS multiple'); f(ws, (re_, 2), f'=B{re_ - 1}/B{rv2 + 2}', '0.00"x"')
f(ws, (re_ + 1, 1), ('When both are on consolidated earnings (FY21) the two agree within 1%. In FY19, NSE\'s standalone figure sat well below '
                     'MOFSL\'s, which is consistent with a large consolidated premium before 2021. The two methods drift apart by 5-9% in later years, '
                     'so treat this as a rough cross-check.'), italic=True, size=9, color=INK2)

# ================================================================= Method & Sources
ws = ws_x
ws.column_dimensions['A'].width = 150
for k, (txt, bold) in enumerate(CFG['method_lines'], start=1):
    c = ws.cell(row=k, column=1, value=txt)
    c.font = Font(name=BODY, size=11 if bold else 10, bold=bold, color=INK)
    c.alignment = Alignment(wrap_text=True, vertical='top')


# ================================================================= charts
def font_props(size=900, color=INK2, bold=False):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(
        sz=size, b=bold, solidFill=color, latin=DFont(typeface=BRAND))), endParaRPr=CharacterProperties())])


XMIN = (dt.date(2016, 9, 30) - dt.date(1899, 12, 30)).days
XMAX = (dt.date(2026, 9, 30) - dt.date(1899, 12, 30)).days + 1


def style(ch, ymin, ymax, major, numfmt, ytitle):
    ch.y_axis.scaling.min = ymin; ch.y_axis.scaling.max = ymax; ch.y_axis.majorUnit = major
    ch.y_axis.number_format = numfmt
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_GRID, w=9525))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.y_axis.txPr = font_props(); ch.y_axis.delete = False
    ch.y_axis.title = ytitle
    ch.y_axis.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=900, b=False, solidFill=INK2, latin=DFont(typeface=BRAND)))
    ch.x_axis.scaling.min = XMIN; ch.x_axis.scaling.max = XMAX; ch.x_axis.majorUnit = 365.25
    ch.x_axis.number_format = 'mmm yyyy'
    ch.x_axis.majorGridlines = None
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_AXIS, w=9525))
    ch.x_axis.txPr = font_props(); ch.x_axis.delete = False; ch.x_axis.majorTickMark = 'out'
    ch.legend.position = 't'; ch.legend.txPr = font_props(1000, INK)
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    # identical inner plot box on every panel so the date axes line up
    ch.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget='inner', xMode='edge', yMode='edge',
                                                           x=0.085, y=0.19, w=0.87, h=0.62))


def line(ch, ws, xcol, ycol, r0, r1, colour, width, name, dash=None):
    s = Series(Reference(ws, min_col=ycol, min_row=r0, max_row=r1), Reference(ws, min_col=xcol, min_row=r0, max_row=r1), title=None)
    s.tx = SeriesLabel(v=name)
    s.graphicalProperties.line.solidFill = colour
    s.graphicalProperties.line.width = int(width * 12700)
    if dash: s.graphicalProperties.line.dashStyle = dash
    s.smooth = False; s.marker.symbol = 'none'
    ch.series.append(s)
    return s


def label_points(s, idxs, pos_list):
    s.dLbls = DataLabelList(); s.dLbls.showVal = False; s.dLbls.showSerName = False; s.dLbls.showCatName = False
    s.dLbls.showLegendKey = False; s.dLbls.showPercent = False
    for idx, pos in zip(idxs, pos_list):
        lbl = DataLabel(idx=idx, showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False, dLblPos=pos)
        lbl.txPr = font_props(1000, INK, bold=True)
        s.dLbls.dLbl.append(lbl)


def new_chart():
    ch = ScatterChart(); ch.scatterStyle = 'lineMarker'
    return ch


def chart_header(ws, title, sub1, sub2):
    ws.sheet_view.showGridLines = False
    for col in range(1, 23):
        ws.column_dimensions[get_column_letter(col)].width = 7.2
    ws.column_dimensions['A'].width = 2.5
    for r, hgt in {1: 10, 2: 30, 3: 30, 4: 30, 5: 8, 6: 34, 7: 6}.items():
        ws.row_dimensions[r].height = hgt
    ws['B2'] = title; ws['B2'].font = Font(name=BRAND, size=17, bold=True, color=INK)
    for r, txt in ((3, sub1), (4, sub2)):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=19)
        c = ws.cell(row=r, column=2, value=txt)
        c.font = Font(name=BRAND, size=11, color=INK2)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    try:
        logo = XLImage(LOGO); logo.width, logo.height = 58, 58; ws.add_image(logo, 'U2')
    except Exception as e:
        print('logo skipped:', e)


def strip(ws, items):
    for cell, formula, colr in items:
        ws[cell] = formula
        ws[cell].font = Font(name=BRAND, size=12, bold=True, color=INK)
        ws[cell].alignment = Alignment(vertical='center')
        ws[cell].border = Border(top=Side(style='thick', color=colr))


def footer(ws, start_row, lines):
    for k, (txt, ital) in enumerate(lines):
        r = start_row + k
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=22)
        c = ws.cell(row=r, column=2, value=txt)
        c.font = Font(name=BRAND, size=8, color=INK3, italic=ital)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 24 if len(txt) > 190 else 13


def page(ws, last_row):
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'A1:V{last_row}'
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.3


S_ = SUM_ROWS
last_idx = M1 - M0
d_start = D0 + int((px['date'] < mon['date'].iloc[0]).sum())
T = CFG['text']

ws = ws_c1
chart_header(ws, T['c1_title'], T['c1_sub1'], T['c1_sub2'])
L = S_['LATEST']; P = S_['P0']
E0 = S_['S0'] + 3
strip(ws, [
    ('B6', f'="Nifty 50   "&TEXT(Summary!C{P},"#,##0")&"  →  "&TEXT(Summary!C{L},"#,##0")&"   ("&TEXT(Summary!I{P},"0.0")&"x)"', C_PRICE),
    ('I6', f'="Profit per unit   ₹"&TEXT(Summary!F{P},"#,##0")&"*  →  ₹"&TEXT(Summary!F{L},"#,##0")&"   ("&TEXT(Summary!D{E0},"0.0")&"x to "&TEXT(Summary!C{E0},"0.0")&"x)"', C_EPS),
    ('Q6', f'="P/E   "&TEXT(Summary!D{P},"0.0")&"x*  →  "&TEXT(Summary!D{L},"0.0")&"x   ("&TEXT(Summary!D{E0 + 2},"+0%;-0%")&" to "&TEXT(Summary!C{E0 + 2},"+0%;-0%")&")"', C_PE),
])
# reference line for the latest P/E: two points, both formulas
R0 = S_['Q0'] + 7
f(ws_s, (R0, 1), 'Chart helper: latest-P/E reference line (two points)', italic=True, size=9, color=INK2)
f(ws_s, (R0 + 1, 1), f'=Monthly!A{M0}', 'dd-mmm-yyyy'); f(ws_s, (R0 + 1, 2), f'=Monthly!C{M1}', '0.00')
f(ws_s, (R0 + 2, 1), f'=Monthly!A{M1}', 'dd-mmm-yyyy'); f(ws_s, (R0 + 2, 2), f'=Monthly!C{M1}', '0.00')

c_price = new_chart()
line(c_price, ws_d, 1, 2, d_start, D1, C_PRICE, 1.5, 'Nifty 50 (index level, daily close)')
style(c_price, 0, 30000, 5000, '#,##0', 'Nifty 50')
c_price.height, c_price.width = 7.0, 37.5

c_eps = new_chart()
s_e1 = line(c_eps, ws_m, 1, 5, M0, sw_row - 1, C_EPS_SA, 2.0, 'Profit per unit, old method: parent company only (to Feb 2021)')
s_e2 = line(c_eps, ws_m, 1, 5, sw_row, M1, C_EPS, 2.25, 'New method, with subsidiaries (from Mar 2021)')
style(c_eps, 0, 1400, 200, '"₹"#,##0', 'Profit per unit, last 4 quarters (₹)')
label_points(s_e2, [M1 - sw_row], ['t'])
c_eps.height, c_eps.width = 7.0, 37.5

c_pe = new_chart()
s_p1 = line(c_pe, ws_m, 1, 3, M0, sw_row - 1, C_PE_SA, 2.0, 'NSE P/E, old method (to Feb 2021)')
s_p2 = line(c_pe, ws_m, 1, 3, sw_row, M1, C_PE, 2.25, 'New method (from Mar 2021)')
line(c_pe, ws_s, 1, 2, R0 + 1, R0 + 2, C_REF, 1.0, T['ref_name'], dash='dash')
style(c_pe, 10, 45, 5, '"\u2007\u2007\u2007"0"x"', 'NSE P/E (trailing)')
label_points(s_p2, [M1 - sw_row], ['b'])
c_pe.height, c_pe.width = 7.0, 37.5

ws.add_chart(c_price, 'B8')
ws.add_chart(c_eps, 'B23')
ws.add_chart(c_pe, 'B38')
footer(ws, 53, [(T['c1_foot1'], False), (T['c1_foot2'], False), ('', False), (T['disclaimer'], True)])
page(ws, 57)

# fonts for body sheets
for sh in (ws_s, ws_v):
    for row in sh.iter_rows():
        for c in row:
            if c.value is not None and c.font.name in (None, 'Calibri'):
                c.font = Font(name=BODY, size=10, color='000000')

wb.save(OUT)
print('saved', OUT, 'daily', len(px), 'monthly', len(mon), 'switch row', sw_row, 'summary', SUM_ROWS)
