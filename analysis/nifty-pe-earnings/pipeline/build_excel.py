"""Build the Nifty 50 price / forward P/E / forward earnings workbook (all derived numbers are formulas)."""
import datetime as dt
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.label import DataLabelList, DataLabel
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont

SCRATCH = '/tmp/claude-0/-home-user-Blockdeals/79809461-fe1d-550f-bdd0-e82c580f77b9/scratchpad'
OUT = f'{SCRATCH}/build/nifty_pe_price_earnings.xlsx'
LOGO = ('/root/.claude/skills/synced/5df40f86-7c7f-40a9-b3c9-2f700a72c2cf_4c4812c9-5adb-4ea7-8980-b5669d92109d/'
        '1finance-brand-guidelines/assets/1finance-logo-vertical-black-200.png')

# ---- brand tokens (1 Finance) -------------------------------------------------------------
C_PRICE, C_EPS, C_PE = '6D57AA', 'CF6827', '356494'      # purple 700, orange 600, blue 600 (validated)
C_REF, C_GRID, C_AXIS = '8E8C81', 'E4E4E0', 'BCBCB5'     # neutral 500, grid, neutral 300
INK, INK2, INK3 = '272623', '6C6A62', '8E8C81'           # neutral 950 / 700 / 500
BRAND = 'Fira Sans'
BODY = 'Arial'
BLUE_INPUT = '0000FF'

inp = pd.read_csv(f'{SCRATCH}/digitize/series_inputs.csv', parse_dates=['date'])
N = len(inp)
FIRST, LAST = 2, N + 1          # data rows on the Data sheet
BASE_DATE = dt.date(2020, 6, 2)

wb = Workbook()
ws_chart = wb.active; ws_chart.title = 'Chart'
ws_data = wb.create_sheet('Data')
ws_sum = wb.create_sheet('Summary')
ws_val = wb.create_sheet('Validation')
ws_meth = wb.create_sheet('Method & Sources')

thin = Side(style='thin', color='D4D4CF')
hdr_fill = PatternFill('solid', fgColor='F6F6FC')       # purple 50


def hdr(ws, row, labels, widths=None):
    for j, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = Font(name=BODY, bold=True, size=10, color=INK)
        c.fill = hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = Border(bottom=thin)
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


# =========================================================================== Data
hdr(ws_data, 1, [
    'Date', 'Nifty 50 close', 'Forward P/E, digitised (band midpoint)', 'Band: lowest possible P/E',
    'Band: highest possible P/E', 'Implied forward EPS (raw)', 'Forward EPS (21-day median)',
    'Forward P/E (rebuilt = close / forward EPS)', 'Nifty 50 index (base = 100)', 'Forward EPS index (base = 100)',
    'Latest P/E (reference line)', 'Rebuilt P/E inside chart band?'],
    [12, 13, 15, 13, 13, 14, 14, 15, 14, 14, 13, 13])
ws_data.row_dimensions[1].height = 58
ws_data.freeze_panes = 'B2'
for i, r in enumerate(inp.itertuples(index=False)):
    row = FIRST + i
    ws_data.cell(row=row, column=1, value=r.date.date()).number_format = 'dd-mmm-yyyy'
    ws_data.cell(row=row, column=2, value=float(r.close)).number_format = '#,##0.00'
    for col, v in ((3, r.pe_dig), (4, r.pe_band_lo), (5, r.pe_band_hi)):
        c = ws_data.cell(row=row, column=col)
        if pd.notna(v) and abs(v) != float('inf'):
            c.value = round(float(v), 4)
        c.number_format = '0.00'
    for col in (2, 3, 4, 5):
        ws_data.cell(row=row, column=col).font = Font(name=BODY, size=10, color=BLUE_INPUT)
    lo, hi = max(FIRST, row - 10), min(LAST, row + 10)
    ws_data.cell(row=row, column=6, value=f'=IF(C{row}="","",B{row}/C{row})').number_format = '#,##0.0'
    ws_data.cell(row=row, column=7, value=f'=MEDIAN(F{lo}:F{hi})').number_format = '#,##0.0'
    ws_data.cell(row=row, column=8, value=f'=B{row}/G{row}').number_format = '0.0"x"'
    ws_data.cell(row=row, column=9, value=f'=B{row}/Summary!$C$8*100').number_format = '0.0'
    ws_data.cell(row=row, column=10, value=f'=G{row}/Summary!$E$8*100').number_format = '0.0'
    ws_data.cell(row=row, column=11, value=f'=Summary!$D$9').number_format = '0.00'
    ws_data.cell(row=row, column=12,
                 value=f'=IF(OR(D{row}="",E{row}=""),"edge hidden",IF(AND(H{row}>=D{row}-0.02,H{row}<=E{row}+0.02),"yes","no"))')
    for col in range(6, 13):
        ws_data.cell(row=row, column=col).font = Font(name=BODY, size=10, color='000000')
ws_data['C1'].comment = Comment('Read off the Bloomberg chart image: midpoint of the orange stroke at this date. '
                                'Blank where the red annotation circle or dashed line hides an edge of the stroke.', 'analysis')
ws_data['D1'].comment = Comment('The true line centre must lie between these two values (stroke edges minus half the '
                                'stroke width). Blank = that edge is hidden by an annotation.', 'analysis')
ws_data['G1'].comment = Comment('Centred rolling median over 21 trading days of the raw implied EPS. Forward EPS moves '
                                'slowly, so this removes digitisation noise without shifting the trend.', 'analysis')

# =========================================================================== Summary
ws = ws_sum
ws.column_dimensions['A'].width = 34
for col, w in zip('BCDEFGHIJ', [14, 13, 12, 13, 13, 13, 13, 13, 13]):
    ws.column_dimensions[col].width = w
ws['A1'] = 'Nifty 50: price, forward P/E and forward earnings at key dates'
ws['A1'].font = Font(name=BODY, bold=True, size=13, color=INK)
ws['A2'] = 'Blue = input you can change. Everything else is a formula on the Data sheet. Price = forward EPS x forward P/E, so the three multiples in each row reconcile.'
ws['A2'].font = Font(name=BODY, italic=True, size=9, color=INK2)
ws['A4'] = 'Index base date (input)'; ws['B4'] = BASE_DATE
ws['B4'].number_format = 'dd-mmm-yyyy'; ws['B4'].font = Font(name=BODY, size=10, color=BLUE_INPUT, bold=True)
ws['C4'] = 'Early June 2020, when the forward P/E was last at today\'s level (the circled point on the Bloomberg chart).'
ws['C4'].font = Font(name=BODY, size=9, color=INK2)
hdr(ws, 6, ['Point in time', 'Date used', 'Nifty 50', 'Fwd P/E', 'Fwd EPS', 'Price multiple to latest',
            'EPS multiple to latest', 'P/E multiple to latest', 'Price CAGR to latest', 'EPS CAGR to latest'])
ws.row_dimensions[6].height = 30
rows = [
    ('Start of 2020 (pre-COVID)', '=INDEX(Data!$A:$A,MATCH(DATE(2020,1,1),Data!$A:$A,1))'),
    ('Index base: early June 2020', '=INDEX(Data!$A:$A,MATCH($B$4,Data!$A:$A,1))'),
    ('Latest', f'=MAX(Data!$A${FIRST}:$A${LAST})'),
    ('COVID low (lowest P/E)', f'=INDEX(Data!$A:$A,MATCH(MIN(Data!$H${FIRST}:$H${LAST}),Data!$H:$H,0))'),
    ('Peak forward P/E', f'=INDEX(Data!$A:$A,MATCH(MAX(Data!$H${FIRST}:$H${LAST}),Data!$H:$H,0))'),
    ('Mid-2022 low', '=INDEX(Data!$A:$A,MATCH(MIN(Data!$H$600:$H$700),Data!$H:$H,0))'),
    ('Nifty all-time closing high', f'=INDEX(Data!$A:$A,MATCH(MAX(Data!$B${FIRST}:$B${LAST}),Data!$B:$B,0))'),
    ('End of 2025', '=INDEX(Data!$A:$A,MATCH(DATE(2025,12,31),Data!$A:$A,1))'),
]
for k, (label, date_f) in enumerate(rows):
    r = 7 + k
    ws.cell(row=r, column=1, value=label).font = Font(name=BODY, size=10, color=INK, bold=(label == 'Latest'))
    ws.cell(row=r, column=2, value=date_f).number_format = 'dd-mmm-yyyy'
    m = f'MATCH($B{r},Data!$A:$A,0)'
    ws.cell(row=r, column=3, value=f'=INDEX(Data!$B:$B,{m})').number_format = '#,##0'
    ws.cell(row=r, column=4, value=f'=INDEX(Data!$H:$H,{m})').number_format = '0.0"x"'
    ws.cell(row=r, column=5, value=f'=INDEX(Data!$G:$G,{m})').number_format = '"₹"#,##0'
    ws.cell(row=r, column=6, value=f'=$C$9/C{r}').number_format = '0.00"x"'
    ws.cell(row=r, column=7, value=f'=$E$9/E{r}').number_format = '0.00"x"'
    ws.cell(row=r, column=8, value=f'=$D$9/D{r}').number_format = '0.00"x"'
    ws.cell(row=r, column=9, value=f'=IF($B$9-B{r}<365,"",(C$9/C{r})^(365.25/($B$9-B{r}))-1)').number_format = '0.0%'
    ws.cell(row=r, column=10, value=f'=IF($B$9-B{r}<365,"",(E$9/E{r})^(365.25/($B$9-B{r}))-1)').number_format = '0.0%'
    for col in range(2, 11):
        ws.cell(row=r, column=col).font = Font(name=BODY, size=10, color='000000')
ws['A9'].fill = PatternFill('solid', fgColor='F6F6FC')
ws['A16'] = 'Check: price multiple = EPS multiple x P/E multiple (base row)'
ws['A16'].font = Font(name=BODY, size=9, italic=True, color=INK2)
ws['B16'] = '=ROUND(F8-G8*H8,6)'
ws['C16'] = '=IF(ABS(B16)<0.0001,"reconciles","CHECK")'
ws['A18'] = 'Latest forward P/E vs earlier lows (x)'
ws['A18'].font = Font(name=BODY, bold=True, size=10, color=INK)
ws['A19'] = 'Latest minus mid-2022 low'; ws['B19'] = '=D9-D12'; ws['B19'].number_format = '+0.00;-0.00'
ws['A20'] = 'Latest minus March-2023 low'
ws['B20'] = '=D9-MIN(Data!$H$840:$H$880)'; ws['B20'].number_format = '+0.00;-0.00'
ws['C19'] = 'Differences this small are inside the digitisation error (about ±0.4x), so treat "cheapest since 2020" as Bloomberg\'s call on its own data.'
ws['C19'].font = Font(name=BODY, size=9, color=INK2)
ws['A22'] = 'Share of days where the rebuilt P/E sits inside the chart band'
ws['B22'] = f'=COUNTIF(Data!$L${FIRST}:$L${LAST},"yes")/(COUNTIF(Data!$L${FIRST}:$L${LAST},"yes")+COUNTIF(Data!$L${FIRST}:$L${LAST},"no"))'
ws['B22'].number_format = '0%'
for r in (19, 20, 22):
    ws.cell(row=r, column=1).font = Font(name=BODY, size=10, color=INK)
    ws.cell(row=r, column=2).font = Font(name=BODY, size=10, color='000000')

# =========================================================================== Validation
ws = ws_val
ws['A1'] = 'A. Digitised forward P/E vs readings published with a Bloomberg attribution'
ws['A1'].font = Font(name=BODY, bold=True, size=12, color=INK)
hdr(ws, 2, ['Date', 'Published fwd P/E', 'Our rebuilt fwd P/E', 'Difference', 'Attribution', 'Source'],
    [13, 14, 14, 12, 34, 70])
pubs = [
    (dt.date(2020, 3, 5), 16.6, 'Bloomberg data, via Financial Express', 'https://prelaunch.financialexpress.com/market/niftys-premium-valuations-take-hit-slip-below-5-yr-average-1890411/'),
    (dt.date(2021, 10, 18), 22.74, 'Bloomberg data, via Financial Express', 'Financial Express (18 May 2022 report citing the 18 Oct 2021 peak)'),
    (dt.date(2022, 5, 18), 17.93, 'Bloomberg data, via Financial Express', 'Financial Express (18 May 2022)'),
    (dt.date(2022, 6, 13), 17.40, 'Bloomberg data, via Financial Express', 'Financial Express (13 Jun 2022)'),
    (dt.date(2024, 9, 4), 21.05, 'Bloomberg data, via Moneycontrol', 'Moneycontrol (4 Sep 2024)'),
    (dt.date(2024, 9, 26), 21.4, 'Bloomberg data, via Paisa Journal', 'Paisa Journal (26 Sep 2024)'),
    (dt.date(2026, 9, 12), 17.6, 'Bloomberg, via ThePrint (approximate)', 'ThePrint republishing Bloomberg (12 Sep 2026)'),
]
for k, (d, v, attr, src) in enumerate(pubs):
    r = 3 + k
    ws.cell(row=r, column=1, value=d).number_format = 'dd-mmm-yyyy'
    ws.cell(row=r, column=2, value=v).number_format = '0.00'
    ws.cell(row=r, column=2).font = Font(name=BODY, size=10, color=BLUE_INPUT)
    ws.cell(row=r, column=3, value=f'=INDEX(Data!$H:$H,MATCH(A{r},Data!$A:$A,1))').number_format = '0.00'
    ws.cell(row=r, column=4, value=f'=C{r}-B{r}').number_format = '+0.00;-0.00'
    ws.cell(row=r, column=5, value=attr); ws.cell(row=r, column=6, value=src)
r_end = 3 + len(pubs) - 1
ws.cell(row=r_end + 1, column=1, value='Mean absolute difference').font = Font(name=BODY, bold=True, size=10)
ws.cell(row=r_end + 1, column=4, value=f'=SUMPRODUCT(ABS(D3:D{r_end}))/COUNT(D3:D{r_end})').number_format = '0.00'
ws.cell(row=r_end + 2, column=1, value=('Excluded: Mint (11 Sep 2020) quoted "Bloomberg data" at 25.5x for 10 Sep 2020; the Bloomberg chart used here '
                                        'shows about 20x that day, so Mint was quoting a different Bloomberg field. Brokerage series (Motilal Oswal, '
                                        'Jefferies, LSEG) run on different earnings estimates and differ by 0.5-2x; they are not used for this check.'))
ws.cell(row=r_end + 2, column=1).font = Font(name=BODY, size=9, italic=True, color=INK2)

r0 = r_end + 5
ws.cell(row=r0, column=1, value='B. Actual Nifty 50 EPS as reported by Motilal Oswal (fiscal years ending March)').font = Font(name=BODY, bold=True, size=12, color=INK)
hdr_row = r0 + 1
for j, lab in enumerate(['Fiscal year', 'EPS (₹)', 'YoY growth', 'Status', 'Note', 'Source'], start=1):
    c = ws.cell(row=hdr_row, column=j, value=lab); c.font = Font(name=BODY, bold=True, size=10, color=INK); c.fill = hdr_fill
eps = [
    ('FY21', 539, 'actual', '+14.2% YoY as reported', 'Mint (Jun 2021), MOFSL'),
    ('FY22', 737, 'actual (interim)', 'From the 4QFY22 review; final print not found', 'Zee Business / PTI (May 2022), MOFSL'),
    ('FY23', 806, 'actual', '', 'MOFSL via BQ Prime (Jul 2023)'),
    ('FY24', 1005, 'actual', '+24% YoY as reported', 'Motilal Oswal (31 May 2024)'),
    ('FY25', 1013, 'actual', '+1% YoY as reported', 'Motilal Oswal (Jun 2025)'),
    ('FY26', 1065, 'actual', '+5% YoY as reported', 'MOFSL via The Adviser Times / ET Now (Jun 2026)'),
    ('FY27E', 1232, 'estimate', 'MOFSL Aug 2026; consensus was 1,357 a year earlier (JM Financial / Bloomberg)', 'Motilal Oswal (Aug 2026)'),
    ('FY28E', 1425, 'estimate', '', 'Motilal Oswal (Aug 2026)'),
]
for k, (fy, v, status, note, src) in enumerate(eps):
    r = hdr_row + 1 + k
    ws.cell(row=r, column=1, value=fy)
    ws.cell(row=r, column=2, value=v).number_format = '#,##0'
    ws.cell(row=r, column=2).font = Font(name=BODY, size=10, color=BLUE_INPUT)
    if k > 0:
        ws.cell(row=r, column=3, value=f'=B{r}/B{r-1}-1').number_format = '0.0%'
    ws.cell(row=r, column=4, value=status); ws.cell(row=r, column=5, value=note); ws.cell(row=r, column=6, value=src)
fy21, fy26 = hdr_row + 1, hdr_row + 6
rr = hdr_row + 1 + len(eps) + 1
ws.cell(row=rr, column=1, value='FY21 to FY26 actual EPS multiple').font = Font(name=BODY, bold=True, size=10)
ws.cell(row=rr, column=2, value=f'=B{fy26}/B{fy21}').number_format = '0.00"x"'
ws.cell(row=rr + 1, column=1, value='FY21 to FY26 actual EPS CAGR').font = Font(name=BODY, bold=True, size=10)
ws.cell(row=rr + 1, column=2, value=f'=(B{fy26}/B{fy21})^(1/5)-1').number_format = '0.0%'
ws.cell(row=rr + 2, column=1, value='Forward EPS multiple, base date to latest (from Summary)').font = Font(name=BODY, bold=True, size=10)
ws.cell(row=rr + 2, column=2, value='=Summary!G8').number_format = '0.00"x"'
ws.cell(row=rr + 5, column=1, value='C. Consensus FY27 EPS and the 12-month forward blend').font = Font(name=BODY, bold=True, size=12, color=INK)
cons = [('Consensus FY27 EPS, mid-2025 (₹)', 1357, 'JM Financial Nifty50 Analyser (Bloomberg consensus), via Moneycontrol / 5paisa, 3 Jun 2026'),
        ('Consensus FY27 EPS, May 2026 (₹)', 1235, 'Same source')]
for k, (lab, v, src) in enumerate(cons):
    r = rr + 6 + k
    ws.cell(row=r, column=1, value=lab)
    ws.cell(row=r, column=2, value=v).number_format = '#,##0'
    ws.cell(row=r, column=2).font = Font(name=BODY, size=10, color=BLUE_INPUT)
    ws.cell(row=r, column=6, value=src)
ws.cell(row=rr + 8, column=1, value='Cut in consensus FY27 EPS')
ws.cell(row=rr + 8, column=2, value=f'=B{rr + 7}/B{rr + 6}-1').number_format = '0.0%'
ws.cell(row=rr + 9, column=1, value='50/50 blend of MOFSL FY27E and FY28E (₹): what a 12-month forward figure looks like in late September')
ws.cell(row=rr + 9, column=2, value=f'=0.5*B{hdr_row + 7}+0.5*B{hdr_row + 8}').number_format = '#,##0'
ws.cell(row=rr + 10, column=1, value='Latest implied forward EPS on the Data sheet (₹)')
ws.cell(row=rr + 10, column=2, value='=Summary!E9').number_format = '#,##0'
ws.cell(row=rr + 11, column=1, value=('Actual EPS roughly doubled FY21-FY26; the rest of the 2.3x rise in forward EPS is forecast, since the '
                                      'late-September 12-month window is about half FY27 and half FY28.'))
ws.cell(row=rr + 11, column=1).font = Font(name=BODY, size=9, italic=True, color=INK2)
ws.cell(row=rr + 3, column=1, value=('FY20 actual EPS is not cleanly reported: MOFSL\'s FY21 figure implies about ₹472, other houses '
                                     'put FY20 near ₹400, because constituents and consolidation differ. So the comparison starts at FY21.'))
ws.cell(row=rr + 3, column=1).font = Font(name=BODY, size=9, italic=True, color=INK2)
for row in ws.iter_rows(min_row=3, max_row=rr + 2):
    for c in row:
        if c.font is None or c.font.name != BODY or (c.font.color is None):
            pass
ws.column_dimensions['E'].width = 44

# =========================================================================== Method & Sources
ws = ws_meth
ws.column_dimensions['A'].width = 130
lines = [
    ('How this workbook was built', True),
    ('1. Nifty 50 daily closes (31 Oct 2019 to 28 Sep 2026) come from Yahoo Finance (^NSEI, NSE data). Two people-independent transcriptions of every page were diffed; all 1,728 rows matched exactly, and every month-end close matches Yahoo\'s monthly table.', False),
    ('2. The 12-month forward P/E is Bloomberg\'s series from the chart "Indian Stocks Cheapest Since 2020 by Forward P/E" (Source: Bloomberg). Bloomberg does not publish this series openly, so it was digitised from the chart image (4096 x 2223 px).', False),
    ('   Y scale: seven gridlines (12 to 24) fitted linearly, largest residual 0.003x. X scale: seven year ticks fitted linearly, largest residual 0.24 days; the COVID low lands on 22-23 Mar 2020.', False),
    ('   For each trading day the top and bottom edge of the orange stroke are read. Half the stroke width (4.25 px, about 0.03x) is taken off each edge, which leaves the range the true value must lie in.', False),
    ('3. Implied forward EPS = Nifty close / digitised forward P/E, then a centred 21-trading-day median (forward EPS moves slowly; the median removes reading noise).', False),
    ('4. Rebuilt forward P/E = Nifty close / smoothed forward EPS. This puts back the real day-to-day moves from actual prices. It sits inside the chart band on about 86% of days; when outside, the median miss is 0.03x and the worst is 0.26x.', False),
    ('5. Against seven readings published with a Bloomberg attribution (Validation sheet), the rebuilt series differs by -0.36x to +0.34x. Treat every P/E here as accurate to about ±0.4x.', False),
    ('6. The last Bloomberg point is 28 Sep 2026 (the final downward tick of the line, the day Nifty fell 1.6%). Nothing is extrapolated past the chart.', False),
    ('7. Indices: price and forward EPS are rebased to 100 on the base date on the Summary sheet (2 Jun 2020). Change that date and the chart and table update.', False),
    ('', False),
    ('What the numbers are, and what they are not', True),
    ('Forward EPS is analysts\' consensus estimate of the next 12 months\' earnings, as Bloomberg aggregates it. It is not reported profit. As a cross-check, Motilal Oswal\'s tally of actual Nifty EPS rose from ₹539 (FY21) to ₹1,065 (FY26), about 2x, in line with the forward series.', False),
    ('Forward P/E readings differ by provider (Bloomberg, LSEG, Motilal Oswal, Jefferies) because each uses its own estimates. Compare levels only within one series.', False),
    ('Nifty 50 here is the price index, without dividends.', False),
    ('', False),
    ('Sources', True),
    ('Bloomberg chart supplied by the user: "Indian Stocks Cheapest Since 2020 by Forward P/E", Nifty 50 12-month forward P/E ratio.', False),
    ('Yahoo Finance, NIFTY 50 (^NSEI) historical data: https://finance.yahoo.com/quote/%5ENSEI/history/', False),
    ('Financial Express, 6 Mar 2020: "Nifty\'s premium valuations take hit, slip below 5-yr average" (Bloomberg data, 16.6x).', False),
    ('Financial Express, May and Jun 2022 reports citing Bloomberg forward P/E (22.74x peak on 18 Oct 2021; 17.93x; 17.40x).', False),
    ('Moneycontrol, 4 Sep 2024 (21.05x, Bloomberg); Paisa Journal, 26 Sep 2024 (21.4x, Bloomberg); ThePrint/Bloomberg, 12 Sep 2026 (~17.6x).', False),
    ('Motilal Oswal Financial Services, India Strategy and Nifty earnings reviews (FY21-FY26 EPS; FY27E/FY28E).', False),
    ('Business Standard / CNBC-TV18, 28 Sep 2026 (Nifty close 22,780.25); Indian Express, 2 Jan 2026 (all-time closing high 26,328.55).', False),
    ('Business Today / PTI, Sep 2026 (FPI net equity outflow of about ₹2.45 lakh crore in calendar 2026 to date).', False),
]
for k, (txt, bold) in enumerate(lines, start=1):
    c = ws.cell(row=k, column=1, value=txt)
    c.font = Font(name=BODY, size=11 if bold else 10, bold=bold, color=INK)
    c.alignment = Alignment(wrap_text=True, vertical='top')

# =========================================================================== Chart sheet
ws = ws_chart
ws.sheet_view.showGridLines = False
for col in range(1, 21):
    ws.column_dimensions[get_column_letter(col)].width = 7.2
ws.column_dimensions['A'].width = 2.5
heights = {1: 10, 2: 30, 3: 18, 4: 18, 5: 8, 6: 34, 7: 6}
for r, hgt in heights.items():
    ws.row_dimensions[r].height = hgt

ws['B2'] = 'Since June 2020, the Nifty and its expected earnings have both grown about 2.3 times'
ws['B2'].font = Font(name=BRAND, size=20, bold=True, color=INK)
ws['B3'] = ('That is why the forward P/E is back near 17x, the level Bloomberg flagged as the cheapest since 2020. '
            'Top: price and expected earnings, indexed to 100 on 2 Jun 2020.')
ws['B3'].font = Font(name=BRAND, size=11, color=INK2)
ws['B4'] = 'Bottom: the 12-month forward P/E, which is price divided by expected earnings.'
ws['B4'].font = Font(name=BRAND, size=11, color=INK2)

# key-number strip (formulas)
strip = [
    ('B6', '="Nifty 50   "&TEXT(Summary!C8,"#,##0")&"  →  "&TEXT(Summary!C9,"#,##0")&"   ("&TEXT(Summary!F8,"0.0")&"x)"', C_PRICE),
    ('H6', '="Forward EPS   ₹"&TEXT(Summary!E8,"#,##0")&"  →  ₹"&TEXT(Summary!E9,"#,##0")&"   ("&TEXT(Summary!G8,"0.0")&"x)"', C_EPS),
    ('N6', '="Forward P/E   "&TEXT(Summary!D8,"0.0")&"x  →  "&TEXT(Summary!D9,"0.0")&"x"', C_PE),
]
for cell, f, colr in strip:
    ws[cell] = f
    ws[cell].font = Font(name=BRAND, size=12, bold=True, color=INK)
    ws[cell].alignment = Alignment(vertical='center')
    # coloured key bar in the cell's top border keeps identity out of the text colour
    ws[cell].border = Border(top=Side(style='thick', color=colr))

try:
    logo = XLImage(LOGO); logo.width, logo.height = 58, 58
    ws.add_image(logo, 'S2')
except Exception as e:
    print('logo skipped:', e)


def font_props(size=900, color=INK2, bold=False):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(
        sz=size, b=bold, solidFill=color, latin=DFont(typeface=BRAND))), endParaRPr=CharacterProperties())])


def style_axes(ch, ymin, ymax, major, numfmt):
    ch.y_axis.scaling.min = ymin; ch.y_axis.scaling.max = ymax; ch.y_axis.majorUnit = major
    ch.y_axis.number_format = numfmt
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_GRID, w=9525))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.y_axis.txPr = font_props()
    ch.y_axis.delete = False
    ch.x_axis.number_format = 'yyyy'
    ch.x_axis.majorTimeUnit = 'years'; ch.x_axis.majorUnit = 1
    ch.x_axis.baseTimeUnit = 'days'
    ch.x_axis.majorGridlines = None
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_AXIS, w=9525))
    ch.x_axis.txPr = font_props()
    ch.x_axis.delete = False
    ch.x_axis.majorTickMark = 'out'
    ch.legend.position = 't'
    ch.legend.txPr = font_props(1000, INK)
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))


def add_line(ch, col, colour, width_pt, name, dash=None):
    ref = Reference(ws_data, min_col=col, min_row=CHART_FIRST, max_row=LAST)
    s = Series(ref, title=None)
    s.tx = SeriesLabel(v=name)
    s.graphicalProperties.line.solidFill = colour
    s.graphicalProperties.line.width = int(width_pt * 12700)
    if dash:
        s.graphicalProperties.line.dashStyle = dash
    s.smooth = False
    s.marker.symbol = 'none'
    ch.series.append(s)
    return s


CHART_FIRST = FIRST + int((inp['date'] < '2020-01-01').sum())          # first trading day of 2020
BASE_IDX = int((inp['date'] < pd.Timestamp(BASE_DATE)).sum()) - (CHART_FIRST - FIRST)   # point index of base date in chart
LAST_IDX = LAST - CHART_FIRST
dates = Reference(ws_data, min_col=1, min_row=CHART_FIRST, max_row=LAST)

latest_pe = float((inp['close'] / (inp['close'] / inp['pe_dig']).rolling(21, center=True, min_periods=1).median()).iloc[-1])
top = LineChart()
top.y_axis.crossAx = 500
top.x_axis = DateAxis(crossAx=100)
add_line(top, 9, C_PRICE, 1.75, 'Nifty 50 (price)')
add_line(top, 10, C_EPS, 2.25, 'Nifty 50 expected earnings (12-month forward EPS)')
top.set_categories(dates)
style_axes(top, 50, 275, 50, '0')
top.y_axis.title = 'Index, 2 Jun 2020 = 100'
top.y_axis.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=900, b=False, solidFill=INK2, latin=DFont(typeface=BRAND)))
top.height, top.width = 9.2, 27.2
top.display_blanks = 'gap'

bot = LineChart()
bot.y_axis.crossAx = 500
bot.x_axis = DateAxis(crossAx=100)
pe_series = add_line(bot, 8, C_PE, 1.75, '12-month forward P/E')
add_line(bot, 11, C_REF, 1.0, f'Latest level ({latest_pe:.1f}x)', dash='dash')
bot.set_categories(dates)
style_axes(bot, 10, 24, 2, '0"x"')
bot.y_axis.title = 'Forward P/E (x)'
bot.y_axis.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=900, b=False, solidFill=INK2, latin=DFont(typeface=BRAND)))
pe_series.dLbls = DataLabelList()
pe_series.dLbls.showVal = False; pe_series.dLbls.showSerName = False; pe_series.dLbls.showCatName = False
pe_series.dLbls.showLegendKey = False; pe_series.dLbls.showPercent = False
for idx, pos in ((BASE_IDX, 'r'), (LAST_IDX, 'b')):
    lbl = DataLabel(idx=idx, showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False,
                    dLblPos=pos)
    lbl.txPr = font_props(1000, INK, bold=True)
    pe_series.dLbls.dLbl.append(lbl)
bot.height, bot.width = 6.4, 27.2

ws.add_chart(top, 'B8')
ws.add_chart(bot, 'B27')

foot = [
    (41, ('Source: Bloomberg (Nifty 50 12-month forward P/E, digitised from Bloomberg\'s chart; accurate to about ±0.4x), '
          'NSE via Yahoo Finance (Nifty 50 closes). Forward EPS = Nifty close ÷ forward P/E, 21-day median. Data to 28 Sep 2026.'), 8, INK3, False),
    (42, ('Expected earnings are analyst forecasts for the next 12 months. FY27 consensus fell from ₹1,357 to ₹1,235 in the year to May 2026 '
          '(JM Financial, Bloomberg consensus). Actual reported Nifty EPS rose from ₹539 (FY21) to ₹1,065 (FY26), per Motilal Oswal.'), 8, INK3, False),
    (44, '1 Finance | SEBI Registered Investment Advisor. This content is for information and educational purposes only. '
         'Please consult a Qualified Financial Advisor before making investment decisions.', 8, INK3, True),
]
for r, txt, size, colr, ital in foot:
    c = ws.cell(row=r, column=2, value=txt)
    c.font = Font(name=BRAND, size=size, color=colr, italic=ital)
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = 'A1:T45'
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.3

# fonts for body sheets
for sh in (ws_sum, ws_val):
    for row in sh.iter_rows():
        for c in row:
            if c.value is not None and c.font.name in (None, 'Calibri'):
                c.font = Font(name=BODY, size=10, color='000000')

wb.save(OUT)
print('saved', OUT, 'rows', N)
